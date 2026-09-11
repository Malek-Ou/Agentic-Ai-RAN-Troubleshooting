from pathlib import Path
from tqdm import tqdm
from collections import Counter
import openpyxl                        
from text_splitter import split_documents   # ← import

from langchain_core.documents import Document
from langchain_community.document_loaders import (
    PyMuPDFLoader,
    CSVLoader,
    Docx2txtLoader,
)
import chardet
from config import (
    INGEST_START,
    INGEST_END
)

#  1. Fonctions _load_xxx  (définies EN PREMIER)
# ─────────────────────────────────────────────

def _load_pdf(file_path: Path) -> list[Document]:
    loader = PyMuPDFLoader(str(file_path))
    return loader.load()
def _load_excel(file_path: Path) -> list[Document]:
    """
    Lit un fichier Excel avec openpyxl (pas besoin de 'unstructured').
    Chaque ligne non-vide devient un Document séparé.
    """
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    docs = []

    SKIP_KEYWORDS = {"cover", "about", "change", "changelog", "bom",
                     "spare", "compatibility", "matrix", "personal",
                     "contact", "content", "description"}

    for sheet_name in wb.sheetnames:
        sheet_lower = sheet_name.lower().strip()
        if any(kw in sheet_lower for kw in SKIP_KEYWORDS):
            print(f"[loader] Sheet ignorée : {sheet_name}")
            continue

        ws = wb[sheet_name]
        headers = None

        for row in ws.iter_rows(values_only=True):
            if headers is None:
                headers = [str(h).strip() if h is not None else f"col_{i}"
                           for i, h in enumerate(row)]
                continue

            if all(cell is None for cell in row):
                continue

            content = " | ".join(
                f"{headers[i]}: {str(cell).strip()}"
                for i, cell in enumerate(row)
                if cell is not None and i < len(headers)
            )

            if not content.strip():
                continue

            doc = Document(
                page_content=content,
                metadata={
                    "source":     str(file_path),
                    "sheet_name": sheet_name,
                    "file_name":  file_path.name,
                }
            )
            docs.append(doc)

    wb.close()
    print(f"[loader] Excel '{file_path.name}': {len(docs)} lignes chargées")
    return docs


def _detect_encoding(file_path: Path) -> str:
    """Détecte automatiquement l'encodage d'un fichier texte."""
    with open(file_path, "rb") as f:
        raw = f.read(10_000)          # lire les 10 premiers Ko suffit
    result = chardet.detect(raw)
    encoding = result.get("encoding") or "utf-8"
    confidence = result.get("confidence", 0)
    print(f"[loader] Encodage détecté pour {file_path.name}: "
          f"{encoding} (confiance: {confidence:.0%})")
    return encoding


def _load_csv(file_path: Path) -> list[Document]:
    encoding = _detect_encoding(file_path)
    try:
        loader = CSVLoader(
            file_path=str(file_path),
            csv_args={"delimiter": ","},
            encoding=encoding,
        )
        return loader.load()
    except (UnicodeDecodeError, Exception):
        # Fallback chaîne : utf-8-sig → latin-1 → cp1252
        for fallback in ("utf-8-sig", "latin-1", "cp1252"):
            try:
                print(f"[loader] Retry encodage {fallback} pour {file_path.name}")
                loader = CSVLoader(
                    file_path=str(file_path),
                    csv_args={"delimiter": ","},
                    encoding=fallback,
                )
                return loader.load()
            except Exception:
                continue
        raise RuntimeError(f"Impossible de lire {file_path.name} : encodage inconnu")


def _load_word(file_path: Path) -> list[Document]:
    loader = Docx2txtLoader(str(file_path))
    return loader.load()


# ─────────────────────────────────────────────
#  2. Mapping EXTENSION_LOADERS  (défini APRÈS les fonctions)
# ─────────────────────────────────────────────

EXTENSION_LOADERS = {
    ".pdf":  _load_pdf,
    ".xlsx": _load_excel,
    ".xls":  _load_excel,
    ".csv":  _load_csv,
    ".docx": _load_word,
    ".doc":  _load_word,
}

SUPPORTED_EXTENSIONS = set(EXTENSION_LOADERS.keys())

# Seuil de filtrage adapté par type de fichier
MIN_CONTENT_LENGTH: dict[str, int] = {# On ne veut pas que l'IA perde du temps avec des pages vides ou des fichiers corrompus.
    "pdf":  80,   # pages riches en texte
    "docx": 80,
    "doc":  80,
    "xlsx": 30,   # lignes tabulaires potentiellement courtes
    "xls":  30,
    "csv":  30,   # ex: "L.Thrp.DL.Cell.Avg, 15.5 Mbps" = ~30 chars
}


# ─────────────────────────────────────────────
#  3. Loader universel
# ─────────────────────────────────────────────

def load_documents(docs_dir: str) -> list[Document]:
    """
    Charge tous les fichiers supportés (PDF, Excel, CSV, Word).
    - Seuil de filtrage adapté par extension
    - Encodage CSV auto-détecté avec fallback
    - Métadonnées enrichies : file_name, file_type, doc_type
    """
    docs_dir  = Path(docs_dir)
    all_files = [
        f for f in sorted(docs_dir.iterdir())
        if f.suffix.lower() in SUPPORTED_EXTENSIONS
    ]

    if not all_files:
        raise FileNotFoundError(
            f"Aucun fichier supporté dans {docs_dir}\n"
            f"Extensions acceptées : {SUPPORTED_EXTENSIONS}"
        )
    
    total_available = len(all_files)
    all_files = all_files[INGEST_START:INGEST_END]
    print(f"[loader] {total_available} fichiers disponibles — "
          f"tranche [{INGEST_START}:{INGEST_END}] → {len(all_files)} fichiers à traiter")

    if not all_files:
        raise ValueError(
            f"Tranche vide : INGEST_START={INGEST_START}, INGEST_END={INGEST_END} "
            f"mais seulement {total_available} fichiers disponibles"
        )
    type_counts = Counter(f.suffix.lower() for f in all_files)
    print(f"[loader] {len(all_files)} fichiers trouvés : {dict(type_counts)}")

    all_docs: list[Document] = []

    for file_path in tqdm(all_files, desc="Chargement documents"):
        ext        = file_path.suffix.lower()
        loader_fn  = EXTENSION_LOADERS[ext]
        min_length = MIN_CONTENT_LENGTH.get(ext.lstrip("."), 80)  # ← seuil adapté

        try:
            docs = loader_fn(file_path)

            for doc in docs:
                if len(doc.page_content.strip()) < min_length:   # ← seuil dynamique
                    continue

                doc.metadata["file_name"] = file_path.name
                doc.metadata["file_type"] = ext.lstrip(".")
                doc.metadata["doc_type"]  = _extract_doc_type(file_path.name)
                all_docs.append(doc)

        except Exception as e:
            print(f"[loader] ERREUR {file_path.name}: {e}")

    print(f"[loader] {len(all_docs)} documents/pages chargés au total")
    return all_docs


# ─────────────────────────────────────────────
#  4. Helper _extract_doc_type : donner une "étiquette métier" à chaque fichier
# ─────────────────────────────────────────────

def _extract_doc_type(filename: str) -> str:
    name = filename.lower()
    if "alarm"         in name or "fault"  in name: return "alarm_guide"
    if "installation"  in name:                      return "installation_guide"
    if "configuration" in name or "config" in name:  return "configuration_guide"
    if "maintenance"   in name:                       return "maintenance_guide"
    if "troubleshoot"  in name:                       return "troubleshooting_guide"
    if "kpi"           in name or "metric" in name:   return "kpi_data"
    if "oss"           in name:                        return "alarm_data"
    return "technical_doc"
if __name__ == "__main__":
    print("--- Démarrage de l'ingestion depuis ./docs_lte ---")
    all_docs = load_documents("./docs_lte")
    print(f"Succès ! {len(all_docs)} segments chargés.")

    # ── Analyse des tailles ──────────────────
    pdf_docs   = [d for d in all_docs if d.metadata["file_type"] == "pdf"]
    excel_docs = [d for d in all_docs if d.metadata["file_type"] == "xlsx"]

    pdf_sizes   = [len(d.page_content) for d in pdf_docs]
    excel_sizes = [len(d.page_content) for d in excel_docs]

    print("\n--- Analyse des tailles ---")
    print(f"PDF   → min:{min(pdf_sizes)} | max:{max(pdf_sizes)} | moy:{sum(pdf_sizes)//len(pdf_sizes)}")
    print(f"Excel → min:{min(excel_sizes)} | max:{max(excel_sizes)} | moy:{sum(excel_sizes)//len(excel_sizes)}")
    chunks = split_documents(all_docs)      # ← chunking

